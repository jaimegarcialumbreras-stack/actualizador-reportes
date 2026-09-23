import io
from datetime import datetime

import pandas as pd
import streamlit as st
import xlrd
from openpyxl import load_workbook

st.set_page_config(page_title="Actualizador de reportes", page_icon="📊")

st.title("📊 Actualizador de reportes")
st.write("Sube Coupons, Promotional Models y el reporte original. Descarga el reporte actualizado.")

st.info(
    "Flujo: Coupons se carga en la hoja Raw data y Promotional Models se carga en la hoja S.O. "
    "No se utilizan archivos CSV."
)

coupons_file = st.file_uploader("Coupons (.xls o .xlsx)", type=["xls", "xlsx"])
promo_file = st.file_uploader("Promotional Models (.xls o .xlsx)", type=["xls", "xlsx"])
report_file = st.file_uploader("Reporte original (.xlsx)", type=["xlsx"])


def is_xls(uploaded_file):
    return uploaded_file.name.lower().endswith(".xls")


def read_excel_file(uploaded_file, preferred_sheet_text=None):
    """Lee un Excel .xls o .xlsx y devuelve sus datos y la hoja seleccionada."""
    uploaded_file.seek(0)

    if is_xls(uploaded_file):
        book = xlrd.open_workbook(
            file_contents=uploaded_file.read(),
            ignore_workbook_corruption=True,
        )
        sheet_names = book.sheet_names()

        if not sheet_names:
            raise ValueError(f'El archivo "{uploaded_file.name}" no contiene hojas.')

        selected_sheet = sheet_names[0]
        if preferred_sheet_text:
            selected_sheet = next(
                (
                    name
                    for name in sheet_names
                    if preferred_sheet_text.casefold() in name.casefold()
                ),
                selected_sheet,
            )

        sheet = book.sheet_by_index(sheet_names.index(selected_sheet))
        rows = [sheet.row_values(row_number) for row_number in range(sheet.nrows)]

        if not rows:
            raise ValueError(
                f'La hoja "{selected_sheet}" del archivo "{uploaded_file.name}" está vacía.'
            )

        headers = [str(value).strip() for value in rows[0]]
        dataframe = pd.DataFrame(rows[1:], columns=headers)
        return dataframe, selected_sheet

    uploaded_file.seek(0)
    excel = pd.ExcelFile(uploaded_file, engine="openpyxl")
    sheet_names = excel.sheet_names

    if not sheet_names:
        raise ValueError(f'El archivo "{uploaded_file.name}" no contiene hojas.')

    selected_sheet = sheet_names[0]
    if preferred_sheet_text:
        selected_sheet = next(
            (
                name
                for name in sheet_names
                if preferred_sheet_text.casefold() in name.casefold()
            ),
            selected_sheet,
        )

    uploaded_file.seek(0)
    dataframe = pd.read_excel(
        uploaded_file,
        sheet_name=selected_sheet,
        engine="openpyxl",
    )
    return dataframe, selected_sheet


def find_column(dataframe, expected_name):
    expected = expected_name.strip().casefold()
    for column in dataframe.columns:
        if str(column).strip().casefold() == expected:
            return column
    return None


def iso_week_number(value):
    if pd.isna(value) or str(value).strip() == "":
        return ""

    try:
        if isinstance(value, datetime):
            parsed_date = value
        else:
            parsed_date = pd.to_datetime(value, dayfirst=True, errors="coerce")

        if pd.isna(parsed_date):
            return ""

        return int(parsed_date.isocalendar().week)
    except (TypeError, ValueError, AttributeError):
        return ""


def prepare_coupons(dataframe):
    dataframe = dataframe.copy()
    date_column = find_column(dataframe, "Fecha Compra")

    if date_column is None:
        columns = ", ".join(str(column) for column in dataframe.columns)
        raise ValueError(
            'No se encontró la columna "Fecha Compra" en Coupons. '
            f"Columnas detectadas: {columns}"
        )

    dataframe["Num de SEM"] = dataframe[date_column].apply(iso_week_number)

    # Columna AC: número 29 para Excel, índice 28 para listas Python.
    ordered_columns = [column for column in dataframe.columns if column != "Num de SEM"]
    ordered_columns.insert(min(28, len(ordered_columns)), "Num de SEM")
    return dataframe[ordered_columns]


def value_for_excel(value):
    if pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    return value


def replace_sheet_data(workbook, sheet_name, dataframe):
    """Borra todos los datos de una hoja existente y escribe el dataframe."""
    if sheet_name not in workbook.sheetnames:
        available = ", ".join(workbook.sheetnames)
        raise ValueError(
            f'No se encontró la hoja "{sheet_name}" en el reporte original. '
            f"Hojas disponibles: {available}"
        )

    worksheet = workbook[sheet_name]

    if worksheet.max_row:
        worksheet.delete_rows(1, worksheet.max_row)

    for column_number, column_name in enumerate(dataframe.columns, start=1):
        worksheet.cell(row=1, column=column_number, value=str(column_name))

    for row_number, row_values in enumerate(
        dataframe.itertuples(index=False, name=None),
        start=2,
    ):
        for column_number, value in enumerate(row_values, start=1):
            worksheet.cell(
                row=row_number,
                column=column_number,
                value=value_for_excel(value),
            )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions


if st.button("Procesar y descargar reporte", type="primary"):
    missing_files = []
    if coupons_file is None:
        missing_files.append("Coupons")
    if promo_file is None:
        missing_files.append("Promotional Models")
    if report_file is None:
        missing_files.append("Reporte original")

    if missing_files:
        st.error("Faltan estos archivos: " + ", ".join(missing_files))
        st.stop()

    try:
        with st.spinner("Procesando archivos..."):
            coupons_df, coupons_sheet = read_excel_file(
                coupons_file,
                preferred_sheet_text="coupon",
            )
            coupons_df = prepare_coupons(coupons_df)

            promo_df, promo_sheet = read_excel_file(promo_file)

            report_file.seek(0)
            workbook = load_workbook(report_file)

            # Destinos solicitados:
            # Coupons -> Raw data; Promotional Models -> S.O.
            replace_sheet_data(workbook, "Raw data", coupons_df)
            replace_sheet_data(workbook, "S.O.", promo_df)

            output = io.BytesIO()
            workbook.save(output)
            output.seek(0)

        st.success("Reporte actualizado correctamente.")
        st.caption(
            f"Coupons: hoja usada “{coupons_sheet}” → Raw data. "
            f"Promotional Models: hoja usada “{promo_sheet}” → S.O."
        )

        left, right = st.columns(2)
        left.metric("Filas cargadas en Raw data", len(coupons_df))
        right.metric("Filas cargadas en S.O.", len(promo_df))

        st.download_button(
            label="Descargar reporte actualizado (.xlsx)",
            data=output.getvalue(),
            file_name="reporte_actualizado.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
        )

    except Exception as error:
        st.error(f"Error al procesar los archivos: {error}")
        st.exception(error)
